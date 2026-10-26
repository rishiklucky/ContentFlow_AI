import sys
import io
import uuid
from pathlib import Path
import pytest
import fitz  # PyMuPDF
import docx

# Add backend directory to path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastapi.testclient import TestClient
from app.main import app
from app.core.database import SessionLocal
from app.models.user import User
from app.models.source import Source
from app.core.security import get_password_hash, create_access_token
from app.services.ssrf import validate_and_resolve_url

client = TestClient(app)


@pytest.fixture(scope="module")
def setup_users():
    db = SessionLocal()
    u1_id = str(uuid.uuid4())
    u2_id = str(uuid.uuid4())

    user1 = User(
        id=u1_id,
        email=f"user1_{u1_id[:8]}@contentflow.ai",
        full_name="User One",
        hashed_password=get_password_hash("Password123!"),
        is_active=True,
    )
    user2 = User(
        id=u2_id,
        email=f"user2_{u2_id[:8]}@contentflow.ai",
        full_name="User Two",
        hashed_password=get_password_hash("Password123!"),
        is_active=True,
    )
    db.add(user1)
    db.add(user2)
    db.commit()

    token1 = create_access_token(subject=u1_id)
    token2 = create_access_token(subject=u2_id)

    yield {
        "user1": user1,
        "token1": token1,
        "user2": user2,
        "token2": token2,
    }

    # Teardown
    try:
        db.query(Source).filter(Source.user_id.in_([u1_id, u2_id])).delete(synchronize_session=False)
        db.query(User).filter(User.id.in_([u1_id, u2_id])).delete(synchronize_session=False)
        db.commit()
    finally:
        db.close()


def test_pasted_text_ingestion(setup_users):
    token = setup_users["token1"]
    headers = {"Authorization": f"Bearer {token}"}

    payload = {
        "title": "Quarterly Financial Analysis",
        "text": "ContentFlow AI transforms source documents into executive summaries, LinkedIn posts, and press releases. Factual accuracy is preserved throughout.",
    }
    res = client.post("/api/v1/sources/text", json=payload, headers=headers)
    assert res.status_code == 201, f"Failed to ingest text: {res.text}"
    data = res.json()
    assert data["title"] == payload["title"]
    assert data["source_type"] == "pasted_text"
    assert data["word_count"] > 10
    assert data["character_count"] > 50
    assert data["sha256_hash"] is not None


def test_pdf_upload_and_extraction(setup_users):
    token = setup_users["token1"]
    headers = {"Authorization": f"Bearer {token}"}

    # Generate a small in-memory PDF using PyMuPDF
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "ContentFlow AI Phase 3 PDF Test Document.\n\nSection 1: Automated Ingestion.\n\nSection 2: High Reliability.")
    pdf_bytes = doc.tobytes()
    doc.close()

    files = {
        "file": ("test_report.pdf", io.BytesIO(pdf_bytes), "application/pdf")
    }
    data = {"title": "Automated PDF Report"}

    res = client.post("/api/v1/sources/upload", files=files, data=data, headers=headers)
    assert res.status_code == 201, f"PDF upload failed: {res.text}"
    source_data = res.json()
    assert source_data["file_type"] == "pdf"
    assert "Section 1: Automated Ingestion" in source_data["extracted_text"]
    assert source_data["meta_info"]["page_count"] == 1
    assert source_data["sha256_hash"] is not None


def test_docx_upload_and_extraction(setup_users):
    token = setup_users["token1"]
    headers = {"Authorization": f"Bearer {token}"}

    # Generate a small in-memory DOCX using python-docx
    doc = docx.Document()
    doc.add_heading("Product Specification", level=1)
    doc.add_paragraph("ContentFlow AI extracts paragraphs and tables seamlessly.")
    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_bytes = docx_io.getvalue()

    files = {
        "file": ("specs.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    }
    res = client.post("/api/v1/sources/upload", files=files, headers=headers)
    assert res.status_code == 201, f"DOCX upload failed: {res.text}"
    source_data = res.json()
    assert source_data["file_type"] == "docx"
    assert "Product Specification" in source_data["extracted_text"]


def test_empty_file_validation(setup_users):
    token = setup_users["token1"]
    headers = {"Authorization": f"Bearer {token}"}

    files = {
        "file": ("empty.txt", io.BytesIO(b""), "text/plain")
    }
    res = client.post("/api/v1/sources/upload", files=files, headers=headers)
    assert res.status_code == 400
    assert "empty" in res.json()["detail"].lower()


def test_unsupported_file_extension(setup_users):
    token = setup_users["token1"]
    headers = {"Authorization": f"Bearer {token}"}

    files = {
        "file": ("binary_exec.exe", io.BytesIO(b"MZ\x90\x00\x03\x00\x00\x00"), "application/octet-stream")
    }
    res = client.post("/api/v1/sources/upload", files=files, headers=headers)
    assert res.status_code == 400
    assert "unsupported" in res.json()["detail"].lower()


def test_ssrf_protection_blocks_internal_addresses():
    # Loopback address
    with pytest.raises(ValueError) as excinfo:
        validate_and_resolve_url("http://127.0.0.1:8000/internal-admin")
    assert "internal or local" in str(excinfo.value).lower() or "restricted" in str(excinfo.value).lower()

    # Localhost
    with pytest.raises(ValueError) as excinfo:
        validate_and_resolve_url("http://localhost:5173")
    assert "internal or local" in str(excinfo.value).lower() or "restricted" in str(excinfo.value).lower()

    # Cloud metadata endpoint
    with pytest.raises(ValueError) as excinfo:
        validate_and_resolve_url("http://169.254.169.254/latest/meta-data")
    assert "restricted" in str(excinfo.value).lower() or "denied" in str(excinfo.value).lower()


def test_user_ownership_isolation(setup_users):
    token1 = setup_users["token1"]
    token2 = setup_users["token2"]

    # User 1 creates a private source
    create_res = client.post(
        "/api/v1/sources/text",
        json={"title": "User 1 Secret Notes", "text": "Confidential internal roadmaps and strategic plans."},
        headers={"Authorization": f"Bearer {token1}"},
    )
    assert create_res.status_code == 201
    source_id = create_res.json()["id"]

    # User 2 tries to access User 1's source -> Must return 404
    get_res = client.get(
        f"/api/v1/sources/{source_id}",
        headers={"Authorization": f"Bearer {token2}"},
    )
    assert get_res.status_code == 404, "User 2 should not be able to view User 1's source"

    # User 2 tries to delete User 1's source -> Must return 404
    del_res = client.delete(
        f"/api/v1/sources/{source_id}",
        headers={"Authorization": f"Bearer {token2}"},
    )
    assert del_res.status_code == 404, "User 2 should not be able to delete User 1's source"

    # User 1 can access and delete their own source
    own_get = client.get(
        f"/api/v1/sources/{source_id}",
        headers={"Authorization": f"Bearer {token1}"},
    )
    assert own_get.status_code == 200
    assert own_get.json()["title"] == "User 1 Secret Notes"

    own_del = client.delete(
        f"/api/v1/sources/{source_id}",
        headers={"Authorization": f"Bearer {token1}"},
    )
    assert own_del.status_code == 200
