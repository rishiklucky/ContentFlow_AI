import io
import os
import re
from typing import Tuple, Dict, Any, List
import fitz  # PyMuPDF
import docx
from PIL import Image

try:
    import pytesseract
    PYTESSERACT_AVAILABLE = True
except ImportError:
    PYTESSERACT_AVAILABLE = False


def calculate_text_metrics(text: str) -> Tuple[int, int, int]:
    """
    Computes character count, word count, and paragraph count.
    """
    clean_text = text.strip()
    if not clean_text:
        return 0, 0, 0

    char_count = len(clean_text)
    word_count = len(re.findall(r"\b\w+\b", clean_text))
    
    # Paragraphs separated by one or more blank lines
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", clean_text) if p.strip()]
    paragraph_count = max(len(paragraphs), 1)

    return char_count, word_count, paragraph_count


def extract_from_txt(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """
    Extracts text from plain text (.txt) with robust encoding detection.
    """
    for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
        try:
            text = file_bytes.decode(encoding).strip()
            return text, {"encoding": encoding}
        except (UnicodeDecodeError, AttributeError):
            continue

    text = file_bytes.decode("utf-8", errors="replace").strip()
    return text, {"encoding": "utf-8-lossy"}


def extract_from_pdf(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """
    Extracts text from PDF documents using PyMuPDF (fitz).
    Detects scanned pages and performs OCR fallback where available.
    """
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as e:
        raise ValueError(f"Corrupted or invalid PDF file: {str(e)}")

    page_texts: List[str] = []
    total_pages = len(doc)
    is_scanned_pdf = False
    ocr_applied = False

    for page_num in range(total_pages):
        page = doc[page_num]
        text = page.get_text("text").strip()

        # If page has virtually no text, it might be a scanned image
        if len(text) < 30 and PYTESSERACT_AVAILABLE:
            try:
                pix = page.get_pixmap(dpi=150)
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                ocr_text = pytesseract.image_to_string(img).strip()
                if len(ocr_text) > len(text):
                    text = ocr_text
                    ocr_applied = True
                    is_scanned_pdf = True
            except Exception:
                # OCR not available or binary not found
                is_scanned_pdf = True

        if text:
            page_texts.append(f"--- Page {page_num + 1} ---\n{text}")

    full_text = "\n\n".join(page_texts).strip()
    
    if not full_text:
        full_text = "[Scanned or image-only PDF document. No machine-readable text was detected.]"
        is_scanned_pdf = True

    metadata = {
        "page_count": total_pages,
        "is_scanned": is_scanned_pdf,
        "ocr_applied": ocr_applied,
        "pdf_metadata": {k: v for k, v in doc.metadata.items() if v} if doc.metadata else {},
    }
    doc.close()

    return full_text, metadata


def extract_from_docx(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """
    Extracts formatted text paragraphs and table data from Microsoft Word (.docx).
    """
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise ValueError(f"Corrupted or invalid Word (.docx) document: {str(e)}")

    parts: List[str] = []

    # Extract paragraphs
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            parts.append(txt)

    # Extract tables
    table_count = len(doc.tables)
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                parts.append(" | ".join(row_cells))

    full_text = "\n\n".join(parts).strip()
    if not full_text:
        full_text = "[Empty Word document with no text content]"

    metadata = {
        "paragraph_count": len(doc.paragraphs),
        "table_count": table_count,
    }

    return full_text, metadata


def extract_from_image(file_bytes: bytes, filename: str) -> Tuple[str, Dict[str, Any]]:
    """
    Extracts text from images (PNG, JPEG, WEBP) using Pillow and Tesseract OCR.
    Handles environments where the Tesseract binary is not installed gracefully.
    """
    try:
        img = Image.open(io.BytesIO(file_bytes))
        width, height = img.size
        img_format = img.format or "UNKNOWN"
    except Exception as e:
        raise ValueError(f"Corrupted or unreadable image file: {str(e)}")

    metadata = {
        "width": width,
        "height": height,
        "format": img_format,
        "mode": img.mode,
        "ocr_available": PYTESSERACT_AVAILABLE,
    }

    extracted_text = ""
    if PYTESSERACT_AVAILABLE:
        try:
            # Convert to RGB if palette/RGBA
            if img.mode not in ("L", "RGB"):
                img = img.convert("RGB")
            extracted_text = pytesseract.image_to_string(img).strip()
            metadata["ocr_applied"] = True
        except Exception as ocr_err:
            metadata["ocr_applied"] = False
            metadata["ocr_note"] = f"OCR engine could not be executed: {str(ocr_err)}"
    else:
        metadata["ocr_applied"] = False
        metadata["ocr_note"] = "Tesseract OCR is not installed on this system."

    if not extracted_text:
        extracted_text = f"[Image File: {filename} ({width}x{height} {img_format}) - No text extracted]"

    return extracted_text, metadata


def extract_content(file_bytes: bytes, filename: str) -> Tuple[str, str, Dict[str, Any]]:
    """
    Detects file format and extracts text and metadata.
    Returns: (file_type, extracted_text, metadata)
    """
    if not file_bytes or len(file_bytes) == 0:
        raise ValueError("Uploaded file is empty (0 bytes).")

    ext = os.path.splitext(filename)[1].lower().lstrip(".")
    if not ext:
        raise ValueError("File has no extension. Please upload a PDF, DOCX, TXT, or Image file.")

    if ext == "pdf":
        text, meta = extract_from_pdf(file_bytes)
        return "pdf", text, meta
    elif ext in ["docx", "doc"]:
        if ext == "doc":
            raise ValueError("Legacy .doc format is not supported. Please convert your file to .docx or .pdf.")
        text, meta = extract_from_docx(file_bytes)
        return "docx", text, meta
    elif ext in ["txt", "md", "csv", "log"]:
        text, meta = extract_from_txt(file_bytes)
        return "txt", text, meta
    elif ext in ["png", "jpg", "jpeg", "webp"]:
        text, meta = extract_from_image(file_bytes, filename)
        return ext, text, meta
    else:
        raise ValueError(
            f"Unsupported file format '.{ext}'. Supported formats are: PDF, DOCX, TXT, PNG, JPG, JPEG, WEBP."
        )
